"""Дрібні файлові та Excel-хелпери (замість зовнішнього пакета vad_utils)."""

import os
import shutil
from pathlib import Path

import openpyxl

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".gif", ".webp", ".tif", ".tiff", ".jfif"}


def read_excel_rows(file_path, sheet_name, min_row=2):
    """Рядки листа (без заголовка) як список списків значень."""
    workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Лист '{sheet_name}' не знайдено в Excel-файлі.")
        sheet = workbook[sheet_name]
        return [list(row) for row in sheet.iter_rows(min_row=min_row, values_only=True)]
    finally:
        workbook.close()


def cell_text(value):
    """Значення комірки як рядок: 12345.0 -> '12345', None -> ''."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def write_lines(file_path, lines):
    with open(file_path, "w", encoding="utf-8") as f:
        for line in lines:
            f.write(f"{line}\n")


def read_lines(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return [line.strip() for line in f if line.strip()]


def ensure_dir(path):
    os.makedirs(path, exist_ok=True)
    return path


def move_file(source, destination):
    """Переміщує файл, перезаписуючи наявний."""
    source, destination = Path(source), Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        destination.unlink()
    shutil.move(str(source), str(destination))


def list_images(directory, extensions=IMAGE_EXTENSIONS):
    directory = Path(directory)
    if not directory.is_dir():
        return []
    return sorted(
        str(p) for p in directory.iterdir()
        if p.is_file() and p.suffix.lower() in extensions
    )
